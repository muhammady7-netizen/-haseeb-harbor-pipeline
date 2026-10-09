#!/usr/bin/env python3
"""Materialise the ORACLE's deliverables into the agent workspace (the golden trajectory is gym calls only)."""
import json, os
from pathlib import Path
sol = Path(os.environ.get("SOLUTION_DIR", "/solution"))
ws = Path(os.environ.get("WORKSPACE", os.environ.get("TH_WORKSPACE_DIR", "/workspace")))
ws.mkdir(parents=True, exist_ok=True)
plan = json.loads((sol / "artifact_plan.json").read_text())
from openpyxl import Workbook
wb = Workbook(); wb.remove(wb.active)
for sheet, rows in plan["xlsx_rows"].items():
    w = wb.create_sheet(title=sheet)
    for row in rows:
        w.append(row)
wb.save(ws / plan["xlsx_path"]); print("wrote", plan["xlsx_path"])
from docx import Document
d = Document(); spec = plan["docx"]
d.add_heading(spec["title"], level=1)
for p in spec["paragraphs"]:
    d.add_paragraph(p)
rows = spec["table"]
t = d.add_table(rows=0, cols=len(rows[0])); t.style = "Table Grid"
for r in rows:
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = str(v)
d.save(ws / plan["docx_path"]); print("wrote", plan["docx_path"])
