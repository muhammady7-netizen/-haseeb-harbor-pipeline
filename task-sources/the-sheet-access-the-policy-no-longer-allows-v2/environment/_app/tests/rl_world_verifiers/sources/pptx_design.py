"""PPTX design facts a client rule can name: geometry, fonts, fills, bullets, OMML.

Every command here reports facts read from the deck's XML; none of them judges.
A task opts in by asking for something a real user would ask for ("keep half an
inch of margin", "no gradients", "titles at least 36 pt") and grading the
matching field with an ordinary comparator (``equals 0``, ``greater_than_equal``).
Nothing in this module is applied to a task that does not name it.

The commands are registered in ``pptx.COMMANDS`` (this module is a helper, not
a source namespace). Only shapes the audience is shown count: hidden slides and
shapes marked ``hidden="1"`` (and the children of a hidden group) are skipped,
and group children are read at any depth with their group transform applied.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator

from pydantic import Field, RootModel, field_validator

from ..models import StrictModel
from ..source_types import SourceCommand, SourceContext, SourceDataError
from .pptx import (
    _effective_run_font,
    _paragraph_property_elements,
    _pptx_theme_root,
    _theme_typeface,
    _validate_pptx_package,
    shape_is_group,
    shape_is_hidden,
    slide_is_hidden,
)

EMU_PER_INCH = 914_400
# Geometry tolerance: 0.01 inch. Coordinates written by different tools round
# differently; an edge-to-edge touch or a one-EMU spill is not a defect.
GEOMETRY_TOLERANCE_EMU = 9_144
MAX_REPORTED_ITEMS = 2_000

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _q(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _local(element: Any) -> str:
    return str(getattr(element, "tag", "")).rsplit("}", 1)[-1]


class PptxPathInput(StrictModel):
    """Input naming one PPTX artifact.

    Attributes:
        path: Workspace-relative PPTX path.
    """

    path: str

    @field_validator("path")
    @classmethod
    def require_pptx_extension(cls, value: str) -> str:
        """Requires the path to target a PPTX file."""
        if Path(value).suffix.lower() != ".pptx":
            raise ValueError("this pptx command requires a .pptx file")
        return value


def _open(path: Path) -> Any:
    from pptx import Presentation

    _validate_pptx_package(path, trusted=False)
    return Presentation(str(path))


def _visible_slides(presentation: Any) -> Iterator[tuple[int, Any]]:
    for index, slide in enumerate(presentation.slides, start=1):
        if not slide_is_hidden(slide):
            yield index, slide


# --------------------------------------------------------------------------
# Shape walking with group transforms
# --------------------------------------------------------------------------

_Transform = tuple[float, float, float, float]  # scale_x, scale_y, offset_x, offset_y
_IDENTITY: _Transform = (1.0, 1.0, 0.0, 0.0)


def _group_transform(group: Any, outer: _Transform) -> _Transform:
    """Compose a group's child-space mapping onto the outer transform."""
    xfrm = None
    properties = group._element.find(_q(_P, "grpSpPr"))
    if properties is not None:
        xfrm = properties.find(_q(_A, "xfrm"))
    if xfrm is None:
        return outer

    def pair(tag: str, a: str, b: str) -> tuple[float, float] | None:
        node = xfrm.find(_q(_A, tag))
        if node is None:
            return None
        try:
            return float(node.get(a, "0")), float(node.get(b, "0"))
        except ValueError:
            return None

    off, ext = pair("off", "x", "y"), pair("ext", "cx", "cy")
    ch_off, ch_ext = pair("chOff", "x", "y"), pair("chExt", "cx", "cy")
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return outer
    sx = ext[0] / ch_ext[0] if ch_ext[0] else 1.0
    sy = ext[1] / ch_ext[1] if ch_ext[1] else 1.0
    # child (x) -> group space: off + (x - chOff) * s ; then the outer mapping.
    osx, osy, oox, ooy = outer
    return (
        osx * sx,
        osy * sy,
        oox + osx * (off[0] - ch_off[0] * sx),
        ooy + osy * (off[1] - ch_off[1] * sy),
    )


def _walk(shapes: Any, transform: _Transform = _IDENTITY, in_group: bool = False,
          group_fill: Any = None) -> Iterator[tuple[Any, _Transform, bool, Any]]:
    """Yield (leaf shape, transform, in_group, nearest group fill element)."""
    for shape in shapes:
        if shape_is_hidden(shape):
            continue
        if shape_is_group(shape):
            properties = shape._element.find(_q(_P, "grpSpPr"))
            fill = _fill_child(properties) if properties is not None else None
            yield from _walk(
                shape.shapes, _group_transform(shape, transform), True,
                fill if fill is not None else group_fill,
            )
        else:
            yield shape, transform, in_group, group_fill


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return (shape.text or "").strip()
    if getattr(shape, "has_table", False):
        return " ".join(
            cell.text.strip() for row in shape.table.rows for cell in row.cells if cell.text.strip()
        )
    return ""


def _shape_kind(shape: Any) -> str:
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    tag = _local(shape._element)
    if tag == "pic":
        return "picture"
    if tag == "cxnSp":
        return "connector"
    if tag == "graphicFrame":
        return "graphic_frame"
    if getattr(shape, "is_placeholder", False):
        return "placeholder"
    if tag == "sp" and _local_text_box(shape):
        return "text_box"
    return "shape"


def _local_text_box(shape: Any) -> bool:
    nv = shape._element.find(_q(_P, "nvSpPr"))
    c_nv = nv.find(_q(_P, "cNvSpPr")) if nv is not None else None
    return c_nv is not None and c_nv.get("txBox") in {"1", "true"}


def _placeholder_type(shape: Any) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    try:
        return str(shape.placeholder_format.type.name or "")
    except (AttributeError, ValueError, KeyError):
        return ""


def _text_role(shape: Any) -> str:
    kind = _placeholder_type(shape)
    if "TITLE" in kind and "SUB" not in kind:
        return "title"
    if kind == "SUBTITLE":
        return "subtitle"
    if kind in {"DATE", "FOOTER", "SLIDE_NUMBER", "HEADER"}:
        return "other"
    if getattr(shape, "has_table", False):
        return "table"
    return "body"


# --------------------------------------------------------------------------
# pptx.inspect_shapes
# --------------------------------------------------------------------------


class PptxInspectShapesInput(PptxPathInput):
    """Input for ``pptx.inspect_shapes``.

    Attributes:
        path: Workspace-relative PPTX path.
        margin_inches: The slide margin a content shape must keep from every
            edge for ``margin_violation``. Default 0.5 inch.
    """

    margin_inches: float = Field(default=0.5, ge=0, le=5)


class PptxShapeGeometry(StrictModel):
    """One visible leaf shape and its box on the slide.

    Attributes:
        slide: 1-based slide index in stored deck order.
        name: Shape name (``cNvPr name``).
        kind: ``text_box``, ``placeholder``, ``shape``, ``picture``, ``table``,
            ``chart``, ``connector`` or ``graphic_frame``.
        placeholder_type: python-pptx placeholder type name, "" otherwise.
        in_group: Whether the shape sits inside a group (box is transformed to
            slide coordinates).
        content: Whether the shape shows content: non-empty text, a table, a
            chart or a picture. Empty placeholders and text-free decoration are
            not content.
        text: First 80 characters of the shape's text.
        left_emu / top_emu / width_emu / height_emu: Box in EMU (914400/inch).
        left_in / top_in / width_in / height_in / right_in / bottom_in: The
            same box in inches, rounded to 3 decimals.
        drawn_box: Whether the box is what the audience sees: a picture, chart,
            table or graphic frame, or a shape with a visible fill or outline.
            A text-only frame (no fill, no outline) is not: its stored size is
            whatever the writer last saved (python-pptx keeps the size it was
            created with; LibreOffice shrinks it to the text), so only its
            left and top edges, where the text starts, are trusted.
        rotated: Whether the shape carries a rotation (the box is unrotated).
        off_slide: The shape extends past a slide edge (0.01-inch tolerance):
            any edge for a drawn box; the left or top edge for a text-only
            frame.
        margin_violation: A content shape on the slide that comes closer than
            ``margin_inches`` to an edge: any edge for a drawn box; the left or
            top edge for a text-only frame.
    """

    slide: int
    name: str
    kind: str
    placeholder_type: str
    in_group: bool
    content: bool
    text: str
    left_emu: int
    top_emu: int
    width_emu: int
    height_emu: int
    left_in: float
    top_in: float
    width_in: float
    height_in: float
    right_in: float
    bottom_in: float
    drawn_box: bool
    rotated: bool
    off_slide: bool
    margin_violation: bool


class PptxOverlap(StrictModel):
    """Two pictures, charts or tables whose boxes intersect.

    Attributes:
        slide: 1-based slide index.
        first / second: Shape names in shape-tree order.
        width_in / height_in: Size of the intersection in inches.
    """

    slide: int
    first: str
    second: str
    width_in: float
    height_in: float


class PptxSlideGeometry(StrictModel):
    """Per-slide geometry counts.

    Attributes:
        index: 1-based slide index.
        off_slide_count / content_off_slide_count / margin_violation_count /
        overlap_count: Counts on this slide.
    """

    index: int
    off_slide_count: int
    content_off_slide_count: int
    margin_violation_count: int
    overlap_count: int


class PptxInspectShapesOutput(StrictModel):
    """Shape geometry on every visible slide.

    Attributes:
        slide_width_in / slide_height_in: Slide size in inches.
        margin_inches: The margin used for ``margin_violation``.
        shapes: Every visible leaf shape with a known box (first 2000).
        slides: Per visible slide counts.
        off_slide_count: Shapes (any kind) extending past a slide edge.
        content_off_slide_count: Content shapes extending past a slide edge.
        margin_violation_count: Content shapes inside the slide but within the
            margin of an edge.
        overlaps: Pairs of pictures, charts or tables whose boxes intersect by
            more than the tolerance in both directions (first 2000). Text
            frames are left out: their stored boxes do not bound their text
            (see ``drawn_box``), so a box overlap is not a text overlap.
        overlap_count: All such pairs.
        unknown_geometry_count: Visible shapes whose box cannot be resolved
            (left out of every count).
    """

    slide_width_in: float
    slide_height_in: float
    margin_inches: float
    shapes: list[PptxShapeGeometry]
    slides: list[PptxSlideGeometry]
    off_slide_count: int
    content_off_slide_count: int
    margin_violation_count: int
    overlaps: list[PptxOverlap]
    overlap_count: int
    unknown_geometry_count: int


def _inches(emu: float) -> float:
    return round(emu / EMU_PER_INCH, 3)


def inspect_pptx_shapes(path: Path, margin_inches: float) -> PptxInspectShapesOutput:
    """Collect geometry facts for ``pptx.inspect_shapes``."""
    presentation = _open(path)
    slide_w = int(presentation.slide_width or 0)
    slide_h = int(presentation.slide_height or 0)
    if slide_w <= 0 or slide_h <= 0:
        raise SourceDataError("PPTX has no slide size")
    margin = margin_inches * EMU_PER_INCH
    tol = GEOMETRY_TOLERANCE_EMU
    shapes_out: list[PptxShapeGeometry] = []
    slides_out: list[PptxSlideGeometry] = []
    overlaps: list[PptxOverlap] = []
    totals = {"off": 0, "content_off": 0, "margin": 0, "overlap": 0, "unknown": 0}
    for index, slide in _visible_slides(presentation):
        boxes: list[tuple[str, float, float, float, float]] = []
        counts = {"off": 0, "content_off": 0, "margin": 0, "overlap": 0}
        for shape, (sx, sy, ox, oy), in_group, group_fill in _walk(slide.shapes):
            try:
                left, top, width, height = shape.left, shape.top, shape.width, shape.height
            except (AttributeError, KeyError, ValueError):
                left = None
            if left is None or top is None or width is None or height is None:
                totals["unknown"] += 1
                continue
            x0, y0 = ox + sx * left, oy + sy * top
            x1, y1 = x0 + sx * width, y0 + sy * height
            x0, x1 = min(x0, x1), max(x0, x1)
            y0, y1 = min(y0, y1), max(y0, y1)
            text = _shape_text(shape)
            kind = _shape_kind(shape)
            exact_kind = kind in {"picture", "table", "chart", "graphic_frame"}
            drawn = exact_kind or _draws_outline_or_fill(shape, group_fill, slide.part)
            content = bool(text) or kind in {"picture", "table", "chart"}
            if drawn:
                off = x0 < -tol or y0 < -tol or x1 > slide_w + tol or y1 > slide_h + tol
                inside_margin = (
                    x0 < margin - tol or y0 < margin - tol
                    or x1 > slide_w - margin + tol or y1 > slide_h - margin + tol
                )
            else:
                off = x0 < -tol or y0 < -tol or x0 > slide_w + tol or y0 > slide_h + tol
                inside_margin = x0 < margin - tol or y0 < margin - tol
            margin_violation = content and not off and inside_margin
            rotated = bool(getattr(shape, "rotation", 0.0) or 0.0)
            counts["off"] += off
            counts["content_off"] += off and content
            counts["margin"] += margin_violation
            name = str(getattr(shape, "name", "") or "")
            if len(shapes_out) < MAX_REPORTED_ITEMS:
                shapes_out.append(PptxShapeGeometry(
                    slide=index, name=name, kind=kind, placeholder_type=_placeholder_type(shape),
                    in_group=in_group, content=content, text=text[:80],
                    left_emu=round(x0), top_emu=round(y0),
                    width_emu=round(x1 - x0), height_emu=round(y1 - y0),
                    left_in=_inches(x0), top_in=_inches(y0),
                    width_in=_inches(x1 - x0), height_in=_inches(y1 - y0),
                    right_in=_inches(x1), bottom_in=_inches(y1), drawn_box=drawn,
                    rotated=rotated, off_slide=off, margin_violation=margin_violation,
                ))
            if kind in {"picture", "table", "chart"}:
                boxes.append((name, x0, y0, x1, y1))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                w = min(a[3], b[3]) - max(a[1], b[1])
                h = min(a[4], b[4]) - max(a[2], b[2])
                if w <= tol or h <= tol:
                    continue
                counts["overlap"] += 1
                if len(overlaps) < MAX_REPORTED_ITEMS:
                    overlaps.append(PptxOverlap(
                        slide=index, first=a[0], second=b[0],
                        width_in=_inches(w), height_in=_inches(h),
                    ))
        slides_out.append(PptxSlideGeometry(
            index=index, off_slide_count=counts["off"],
            content_off_slide_count=counts["content_off"],
            margin_violation_count=counts["margin"], overlap_count=counts["overlap"],
        ))
        for key in counts:
            totals[key] += counts[key]
    return PptxInspectShapesOutput(
        slide_width_in=_inches(slide_w), slide_height_in=_inches(slide_h),
        margin_inches=margin_inches, shapes=shapes_out, slides=slides_out,
        off_slide_count=totals["off"], content_off_slide_count=totals["content_off"],
        margin_violation_count=totals["margin"], overlaps=overlaps,
        overlap_count=totals["overlap"],
        unknown_geometry_count=totals["unknown"],
    )


def _draws_outline_or_fill(shape: Any, group_fill: Any, part: Any) -> bool:
    """Whether the shape draws its box: a visible fill or outline."""
    if _shape_fill_kind(shape, group_fill, part) not in {"none", "unknown"}:
        return True
    properties = shape._element.find(_q(_P, "spPr"))
    line = properties.find(_q(_A, "ln")) if properties is not None else None
    if line is not None:
        if line.find(_q(_A, "noFill")) is not None:
            return False
        if any(line.find(_q(_A, tag)) is not None for tag in ("solidFill", "gradFill", "pattFill")):
            return True
    style = shape._element.find(_q(_P, "style"))
    ref = style.find(_q(_A, "lnRef")) if style is not None else None
    try:
        return ref is not None and int(ref.get("idx", "0")) > 0
    except ValueError:
        return False


class InspectShapes(SourceCommand[PptxInspectShapesInput, PptxInspectShapesOutput]):
    """Report every visible shape's box, off-slide, margin and overlap facts."""

    name = "inspect_shapes"
    input_model = PptxInspectShapesInput
    output_model = PptxInspectShapesOutput

    def run(self, source_input: PptxInspectShapesInput, context: SourceContext) -> PptxInspectShapesOutput:
        resolved = context.resolve_path(source_input.path)
        return inspect_pptx_shapes(resolved, source_input.margin_inches)


# --------------------------------------------------------------------------
# Fonts (shared with bullets)
# --------------------------------------------------------------------------


def _presentation_default_rpr(presentation: Any, level: int) -> Any:
    style = presentation.part._element.find(_q(_P, "defaultTextStyle"))
    if style is None:
        return None
    node = style.find(_q(_A, f"lvl{level + 1}pPr"))
    if node is None:
        return None
    return node.find(_q(_A, "defRPr"))


def _theme_font(part: Any, major: bool) -> str:
    return _theme_typeface(part, "+mj-lt" if major else "+mn-lt")


def _autofit_scale(shape: Any) -> float:
    try:
        body = shape.text_frame._txBody.find(_q(_A, "bodyPr"))
    except AttributeError:
        return 1.0
    if body is None:
        return 1.0
    fit = body.find(_q(_A, "normAutofit"))
    if fit is None:
        return 1.0
    try:
        return int(fit.get("fontScale", "100000")) / 100_000
    except ValueError:
        return 1.0


def _resolved_run_font(presentation: Any, run: Any, paragraph: Any, shape: Any,
                       role: str) -> dict[str, Any]:
    """Effective family and size of a run: the slide/layout/master chain
    (``_effective_run_font``), then the presentation default text style, then
    the theme font (major for titles, minor otherwise) and PowerPoint's 18 pt."""
    font = _effective_run_font(run, paragraph, shape)
    name = font.get("name")
    size = font.get("size")
    level = getattr(paragraph, "level", 0) or 0
    if name is None or size is None:
        default = _presentation_default_rpr(presentation, level)
        if default is not None:
            if name is None:
                latin = default.find(_q(_A, "latin"))
                if latin is not None and latin.get("typeface"):
                    name = _theme_typeface(run.part, latin.get("typeface"))
            if size is None and default.get("sz"):
                try:
                    size = int(default.get("sz"))
                except ValueError:
                    size = None
    if name is None or str(name).startswith("+"):
        name = _theme_font(run.part, role == "title")
    if str(name).startswith("+"):
        name = ""
    size_pt = (size / 100) if isinstance(size, int) else 18.0
    return {"family": str(name or ""), "size_pt": size_pt,
            "bold": bool(font.get("bold")), "italic": bool(font.get("italic"))}


def _iter_text_paragraphs(shape: Any) -> Iterator[tuple[Any, Any]]:
    """(paragraph, owning shape-like for property lookup) for a text frame or table."""
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            yield paragraph, shape
    elif getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    yield paragraph, shape


class PptxRunFont(StrictModel):
    """One text run's effective font.

    Attributes:
        slide: 1-based slide index.
        shape: Shape name.
        role: ``title``, ``subtitle``, ``body``, ``table`` or ``other``
            (date/footer/slide-number placeholders).
        text: First 40 characters of the run.
        family: Effective latin font family (theme tokens resolved).
        size_pt: Effective size in points (inherited sizes resolved; 18 pt when
            nothing in the deck sets one, as PowerPoint does).
        rendered_size_pt: ``size_pt`` times the text frame's shrink-on-overflow
            scale (``normAutofit fontScale``), i.e. what is drawn.
        bold / italic: Effective flags.
    """

    slide: int
    shape: str
    role: str
    text: str
    family: str
    size_pt: float
    rendered_size_pt: float
    bold: bool
    italic: bool


class PptxInspectFontsOutput(StrictModel):
    """Fonts and sizes of every visible text run.

    Size fields are ``None`` when the deck has no run of that role, so a
    ``greater_than_equal`` check on a missing title fails instead of passing.

    Attributes:
        runs: Every non-blank run (first 2000).
        run_count: All non-blank runs.
        families: Distinct effective families, sorted.
        family_count: ``len(families)``.
        title_sizes_pt / body_sizes_pt: Distinct sizes for those roles, sorted.
        min_title_size_pt / max_title_size_pt: Over title runs.
        min_body_size_pt / max_body_size_pt: Over body runs (text boxes, body
            placeholders, shapes; tables, subtitles and footers excluded).
        slide_heading_sizes_pt: Per visible slide, its heading size: the
            largest run in its title placeholder, or, on a slide with no
            titled placeholder (a deck built from text boxes), its largest run
            of any role but ``other``; None for a slide with no text.
        min_slide_heading_size_pt: Smallest of those (None when no slide has
            text). "Titles at least 36 pt" stays gradable on a text-box deck;
            on such a slide a big statistic can stand in for a small title.
        min_size_pt: Over every run except ``other`` (footers, slide numbers).
        min_rendered_size_pt: The same over ``rendered_size_pt``.
        bold_run_count / italic_run_count: Runs with those flags.
    """

    runs: list[PptxRunFont]
    run_count: int
    families: list[str]
    family_count: int
    title_sizes_pt: list[float]
    body_sizes_pt: list[float]
    min_title_size_pt: float | None
    max_title_size_pt: float | None
    min_body_size_pt: float | None
    max_body_size_pt: float | None
    slide_heading_sizes_pt: list[float | None]
    min_slide_heading_size_pt: float | None
    min_size_pt: float | None
    min_rendered_size_pt: float | None
    bold_run_count: int
    italic_run_count: int


def inspect_pptx_fonts(path: Path) -> PptxInspectFontsOutput:
    """Collect effective run fonts for ``pptx.inspect_fonts``."""
    presentation = _open(path)
    runs: list[PptxRunFont] = []
    count = 0
    families: set[str] = set()
    sizes: dict[str, list[float]] = {}
    rendered: list[float] = []
    bold = italic = 0
    headings: list[float | None] = []
    for index, slide in _visible_slides(presentation):
        slide_title: list[float] = []
        slide_any: list[float] = []
        for shape, _, _, _ in _walk(slide.shapes):
            role = _text_role(shape)
            scale = _autofit_scale(shape)
            for paragraph, owner in _iter_text_paragraphs(shape):
                for run in paragraph.runs:
                    if not (run.text or "").strip():
                        continue
                    font = _resolved_run_font(presentation, run, paragraph, owner, role)
                    count += 1
                    if font["family"]:
                        families.add(font["family"])
                    sizes.setdefault(role, []).append(font["size_pt"])
                    if role == "title":
                        slide_title.append(font["size_pt"])
                    elif role != "other":
                        slide_any.append(font["size_pt"])
                    if role != "other":
                        rendered.append(round(font["size_pt"] * scale, 2))
                    bold += font["bold"]
                    italic += font["italic"]
                    if len(runs) < MAX_REPORTED_ITEMS:
                        runs.append(PptxRunFont(
                            slide=index, shape=str(shape.name or ""), role=role,
                            text=run.text.strip()[:40], family=font["family"],
                            size_pt=font["size_pt"],
                            rendered_size_pt=round(font["size_pt"] * scale, 2),
                            bold=font["bold"], italic=font["italic"],
                        ))
        headings.append(max(slide_title) if slide_title else (max(slide_any) if slide_any else None))
    title = sizes.get("title", [])
    body = sizes.get("body", [])
    known_headings = [size for size in headings if size is not None]
    graded = [size for role, values in sizes.items() if role != "other" for size in values]
    return PptxInspectFontsOutput(
        runs=runs, run_count=count, families=sorted(families), family_count=len(families),
        title_sizes_pt=sorted(set(title)), body_sizes_pt=sorted(set(body)),
        min_title_size_pt=min(title) if title else None,
        max_title_size_pt=max(title) if title else None,
        min_body_size_pt=min(body) if body else None,
        max_body_size_pt=max(body) if body else None,
        slide_heading_sizes_pt=headings,
        min_slide_heading_size_pt=min(known_headings) if known_headings else None,
        min_size_pt=min(graded) if graded else None,
        min_rendered_size_pt=min(rendered) if rendered else None,
        bold_run_count=bold, italic_run_count=italic,
    )


class InspectFonts(SourceCommand[PptxPathInput, PptxInspectFontsOutput]):
    """Report effective font families and sizes of every visible run."""

    name = "inspect_fonts"
    input_model = PptxPathInput
    output_model = PptxInspectFontsOutput

    def run(self, source_input: PptxPathInput, context: SourceContext) -> PptxInspectFontsOutput:
        return inspect_pptx_fonts(context.resolve_path(source_input.path))


# --------------------------------------------------------------------------
# pptx.inspect_bullets
# --------------------------------------------------------------------------


class PptxInspectBulletsInput(PptxPathInput):
    """Input for ``pptx.inspect_bullets``.

    Attributes:
        path: Workspace-relative PPTX path.
        large_size_pt: A paragraph whose largest run is at least this size is
            "large" (a headline statistic). Default 40 pt.
    """

    large_size_pt: float = Field(default=40.0, gt=0, le=400)


class PptxParagraphBullet(StrictModel):
    """One non-blank paragraph and its effective bullet.

    Attributes:
        slide: 1-based slide index.
        shape: Shape name.
        role: As in ``inspect_fonts``.
        level: Paragraph indent level (0-based).
        text: First 60 characters.
        bullet: ``char``, ``autonum``, ``picture`` or ``none``, resolved through
            the paragraph, the shape's list style, its layout/master placeholder
            and the master text styles (a body placeholder is bulleted by the
            master even when the paragraph says nothing).
        bullet_char: The character for ``char`` bullets, "" otherwise.
        max_size_pt: Largest effective run size in the paragraph.
    """

    slide: int
    shape: str
    role: str
    level: int
    text: str
    bullet: str
    bullet_char: str
    max_size_pt: float


class PptxInspectBulletsOutput(StrictModel):
    """Bullets on every visible non-blank text-frame paragraph (tables excluded).

    Attributes:
        paragraphs: First 2000 paragraphs.
        paragraph_count: All non-blank paragraphs.
        bulleted_count: Paragraphs with any bullet.
        large_size_pt: Threshold used for ``large_bulleted_count``.
        large_bulleted_count: Bulleted paragraphs whose largest run is at least
            ``large_size_pt`` (a bullet on a big statistic).
        slides_all_bulleted: Visible slides where every body paragraph (two or
            more) is bulleted.
    """

    paragraphs: list[PptxParagraphBullet]
    paragraph_count: int
    bulleted_count: int
    large_size_pt: float
    large_bulleted_count: int
    slides_all_bulleted: list[int]


def _paragraph_bullet(presentation: Any, paragraph: Any, shape: Any) -> tuple[str, str]:
    elements = list(_paragraph_property_elements(paragraph, shape))
    if not getattr(shape, "is_placeholder", False):
        style = presentation.part._element.find(_q(_P, "defaultTextStyle"))
        level = style.find(_q(_A, f"lvl{(paragraph.level or 0) + 1}pPr")) if style is not None else None
        if level is not None:
            elements.append(level)
    for properties in elements:
        for child in properties:
            name = _local(child)
            if name == "buNone":
                return "none", ""
            if name == "buChar":
                return "char", str(child.get("char") or "")
            if name == "buAutoNum":
                return "autonum", ""
            if name == "buBlip":
                return "picture", ""
    return "none", ""


def inspect_pptx_bullets(path: Path, large_size_pt: float) -> PptxInspectBulletsOutput:
    """Collect paragraph bullets for ``pptx.inspect_bullets``."""
    presentation = _open(path)
    out: list[PptxParagraphBullet] = []
    count = bulleted = large = 0
    all_bulleted: list[int] = []
    for index, slide in _visible_slides(presentation):
        body_flags: list[bool] = []
        for shape, _, _, _ in _walk(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            role = _text_role(shape)
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if not text:
                    continue
                kind, char = _paragraph_bullet(presentation, paragraph, shape)
                sizes = [
                    _resolved_run_font(presentation, run, paragraph, shape, role)["size_pt"]
                    for run in paragraph.runs if (run.text or "").strip()
                ]
                max_size = max(sizes) if sizes else 0.0
                count += 1
                is_bulleted = kind != "none"
                bulleted += is_bulleted
                large += is_bulleted and max_size >= large_size_pt
                if role == "body":
                    body_flags.append(is_bulleted)
                if len(out) < MAX_REPORTED_ITEMS:
                    out.append(PptxParagraphBullet(
                        slide=index, shape=str(shape.name or ""), role=role,
                        level=int(paragraph.level or 0), text=text[:60], bullet=kind,
                        bullet_char=char, max_size_pt=max_size,
                    ))
        if len(body_flags) >= 2 and all(body_flags):
            all_bulleted.append(index)
    return PptxInspectBulletsOutput(
        paragraphs=out, paragraph_count=count, bulleted_count=bulleted,
        large_size_pt=large_size_pt, large_bulleted_count=large,
        slides_all_bulleted=all_bulleted,
    )


class InspectBullets(SourceCommand[PptxInspectBulletsInput, PptxInspectBulletsOutput]):
    """Report each visible paragraph's effective bullet and largest run size."""

    name = "inspect_bullets"
    input_model = PptxInspectBulletsInput
    output_model = PptxInspectBulletsOutput

    def run(self, source_input: PptxInspectBulletsInput, context: SourceContext) -> PptxInspectBulletsOutput:
        return inspect_pptx_bullets(context.resolve_path(source_input.path), source_input.large_size_pt)


# --------------------------------------------------------------------------
# pptx.inspect_fills
# --------------------------------------------------------------------------

_FILL_KINDS = {
    "noFill": "none",
    "solidFill": "solid",
    "gradFill": "gradient",
    "blipFill": "picture",
    "pattFill": "pattern",
    "grpFill": "group",
}


def _fill_child(properties: Any) -> Any:
    if properties is None:
        return None
    for child in properties:
        if _local(child) in _FILL_KINDS:
            return child
    return None


def _theme_style_fill(part: Any, idx: int, background: bool) -> str:
    """Fill kind of a theme fill-style reference (``fillRef``/``bgRef``)."""
    if idx == 0:
        return "none"
    root = _pptx_theme_root(part)
    if root is None:
        return "unknown"
    list_name = "bgFillStyleLst" if background and idx >= 1001 else "fillStyleLst"
    position = idx - 1001 if idx >= 1001 else idx - 1
    style_list = root.find(f".//{_q(_A, list_name)}")
    if style_list is None or position < 0 or position >= len(style_list):
        return "unknown"
    return _FILL_KINDS.get(_local(style_list[position]), "unknown")


def _background_fill(slide: Any) -> str:
    owners = [slide]
    try:
        owners.append(slide.slide_layout)
        owners.append(slide.slide_layout.slide_master)
    except (AttributeError, KeyError, ValueError):
        pass
    for owner in owners:
        c_sld = owner._element.find(_q(_P, "cSld"))
        bg = c_sld.find(_q(_P, "bg")) if c_sld is not None else None
        if bg is None:
            continue
        bg_pr = bg.find(_q(_P, "bgPr"))
        if bg_pr is not None:
            child = _fill_child(bg_pr)
            return _FILL_KINDS.get(_local(child), "unknown") if child is not None else "unknown"
        bg_ref = bg.find(_q(_P, "bgRef"))
        if bg_ref is not None:
            try:
                return _theme_style_fill(slide.part, int(bg_ref.get("idx", "0")), True)
            except ValueError:
                return "unknown"
    return "none"


class PptxGradient(StrictModel):
    """One gradient the audience sees.

    Attributes:
        slide: 1-based slide index.
        where: ``background``, ``shape``, ``line``, ``text`` or ``table_cell``.
        shape: Shape name ("" for a background).
    """

    slide: int
    where: str
    shape: str


class PptxSlideFills(StrictModel):
    """Fill facts of one visible slide.

    Attributes:
        index: 1-based slide index.
        background_fill: Effective background (slide, else layout, else master):
            ``solid``, ``gradient``, ``picture``, ``pattern``, ``none`` or
            ``unknown``.
        shape_fills: Count of visible leaf shapes per effective fill kind
            (direct fill, else the group's fill, else the theme style).
    """

    index: int
    background_fill: str
    shape_fills: dict[str, int]


class PptxInspectFillsOutput(StrictModel):
    """Gradient fills anywhere the audience sees them.

    Attributes:
        gradient_count: Gradients across backgrounds, shape fills, outlines,
            text runs and table cells of visible slides.
        gradients: First 2000 of them.
        slides: Per visible slide fill facts.
    """

    gradient_count: int
    gradients: list[PptxGradient]
    slides: list[PptxSlideFills]


def _shape_fill_kind(shape: Any, group_fill: Any, part: Any) -> str:
    if _local(shape._element) == "pic":
        return "picture"
    properties = shape._element.find(_q(_P, "spPr"))
    child = _fill_child(properties)
    if child is not None:
        if _local(child) == "grpFill":
            return _FILL_KINDS.get(_local(group_fill), "none") if group_fill is not None else "none"
        return _FILL_KINDS[_local(child)]
    style = shape._element.find(_q(_P, "style"))
    ref = style.find(_q(_A, "fillRef")) if style is not None else None
    if ref is not None:
        try:
            return _theme_style_fill(part, int(ref.get("idx", "0")), False)
        except ValueError:
            return "unknown"
    return "none"


def inspect_pptx_fills(path: Path) -> PptxInspectFillsOutput:
    """Collect background and shape fills for ``pptx.inspect_fills``."""
    presentation = _open(path)
    gradients: list[PptxGradient] = []
    total = 0
    slides_out: list[PptxSlideFills] = []

    def add(slide_index: int, where: str, name: str) -> None:
        nonlocal total
        total += 1
        if len(gradients) < MAX_REPORTED_ITEMS:
            gradients.append(PptxGradient(slide=slide_index, where=where, shape=name))

    for index, slide in _visible_slides(presentation):
        background = _background_fill(slide)
        if background == "gradient":
            add(index, "background", "")
        fills: dict[str, int] = {}
        for shape, _, _, group_fill in _walk(slide.shapes):
            name = str(shape.name or "")
            kind = _shape_fill_kind(shape, group_fill, slide.part)
            fills[kind] = fills.get(kind, 0) + 1
            if kind == "gradient":
                add(index, "shape", name)
            properties = shape._element.find(_q(_P, "spPr"))
            line = properties.find(_q(_A, "ln")) if properties is not None else None
            if line is not None and line.find(_q(_A, "gradFill")) is not None:
                add(index, "line", name)
            for run_props in shape._element.iter(_q(_A, "rPr")):
                if run_props.find(_q(_A, "gradFill")) is not None:
                    add(index, "text", name)
            for cell_props in shape._element.iter(_q(_A, "tcPr")):
                if cell_props.find(_q(_A, "gradFill")) is not None:
                    add(index, "table_cell", name)
        slides_out.append(PptxSlideFills(index=index, background_fill=background, shape_fills=fills))
    return PptxInspectFillsOutput(gradient_count=total, gradients=gradients, slides=slides_out)


class InspectFills(SourceCommand[PptxPathInput, PptxInspectFillsOutput]):
    """Report gradient fills on backgrounds, shapes, outlines, text and cells."""

    name = "inspect_fills"
    input_model = PptxPathInput
    output_model = PptxInspectFillsOutput

    def run(self, source_input: PptxPathInput, context: SourceContext) -> PptxInspectFillsOutput:
        return inspect_pptx_fills(context.resolve_path(source_input.path))


# --------------------------------------------------------------------------
# pptx.scan_omml
# --------------------------------------------------------------------------


class PptxScanOmmlOutput(StrictModel):
    """Office Math (OMML) equations in the deck.

    Attributes:
        omml_count: ``m:oMath`` elements on visible slides (an equation built
            with PowerPoint's equation editor). Unicode math typed as ordinary
            text is not OMML and does not count.
        slides_with_omml: 1-based indexes of visible slides carrying any.
        hidden_slide_omml_count: The same count on hidden slides.
    """

    omml_count: int
    slides_with_omml: list[int]
    hidden_slide_omml_count: int


def scan_pptx_omml(path: Path) -> PptxScanOmmlOutput:
    """Count OMML equations for ``pptx.scan_omml``."""
    presentation = _open(path)
    visible = hidden = 0
    slides: list[int] = []
    for index, slide in enumerate(presentation.slides, start=1):
        found = sum(1 for _ in slide._element.iter(_q(_M, "oMath")))
        if slide_is_hidden(slide):
            hidden += found
        elif found:
            visible += found
            slides.append(index)
    return PptxScanOmmlOutput(omml_count=visible, slides_with_omml=slides, hidden_slide_omml_count=hidden)


class ScanOmml(SourceCommand[PptxPathInput, PptxScanOmmlOutput]):
    """Count OMML equation elements on visible slides."""

    name = "scan_omml"
    input_model = PptxPathInput
    output_model = PptxScanOmmlOutput

    def run(self, source_input: PptxPathInput, context: SourceContext) -> PptxScanOmmlOutput:
        return scan_pptx_omml(context.resolve_path(source_input.path))



# --------------------------------------------------------------------------
# pptx.read_table
# --------------------------------------------------------------------------


class PptxReadTableInput(PptxPathInput):
    """Input for ``pptx.read_table``.

    Attributes:
        path: Workspace-relative PPTX path.
        header_match: Optional regex tried against each table row rendered as
            ``cell|cell|...``; the first table (visible slides in order, shapes
            in tree order, groups included) with a matching row is read and that
            row is its header. Without it the first table's first row is the
            header (``table_index`` picks another table).
        table_index: Zero-based index among the candidate tables (those whose
            header matches, or all tables without ``header_match``). Default 0.
        all_matching: When True, the data rows of EVERY candidate table are
            concatenated in deck order (a table continued on the next slide
            with the same header row). Default False.
        columns: Optional positional column names (the ``docx.read_table``
            argument). When given they key the grid columns instead of the
            header text, so ``Repository (owner/name)`` and ``Repository``
            read the same; columns beyond the list keep their header text.
    """

    header_match: str | None = Field(default=None, max_length=500)
    table_index: int = Field(default=0, ge=0)
    all_matching: bool = False
    columns: list[str] | None = Field(default=None, max_length=50)

    @field_validator("header_match")
    @classmethod
    def require_valid_regex(cls, value: str | None) -> str | None:
        """The header pattern must compile."""
        import re

        if value is not None:
            try:
                re.compile(value)
            except re.error as exc:
                raise ValueError(f"header_match is not a valid regex: {exc}") from exc
        return value


class PptxReadTableOutput(RootModel[list[dict[str, str]]]):
    """One header-keyed mapping per data row, in table order: the
    ``docx.read_table`` / ``md.read_table`` contract, so ``table_equals``
    grades it. Cell text is whitespace-collapsed. A merged cell's text fills
    every grid slot it spans; a header key that repeats gets its 1-based
    column appended (``Status (3)``); an all-blank data row is skipped."""


def _table_grid(table: Any) -> list[list[str]]:
    rows = list(table.rows)
    width = len(table.columns)
    grid = [["" for _ in range(width)] for _ in rows]
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            grid[r][c] = " ".join((cell.text or "").split())
    # A spanned cell is empty in the XML; give it the origin cell's text.
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            if getattr(cell, "is_merge_origin", False):
                for dr in range(cell.span_height):
                    for dc in range(cell.span_width):
                        if r + dr < len(grid) and c + dc < width and (dr or dc):
                            grid[r + dr][c + dc] = grid[r][c]
    return grid


def read_pptx_table(path: Path, header_match: str | None, table_index: int, all_matching: bool,
                    columns: list[str] | None = None) -> list[dict[str, str]]:
    """Header-keyed rows of the selected slide table(s)."""
    import re

    presentation = _open(path)
    pattern = re.compile(header_match) if header_match else None
    candidates: list[tuple[list[list[str]], int]] = []
    for _, slide in _visible_slides(presentation):
        for shape, _, _, _ in _walk(slide.shapes):
            if not getattr(shape, "has_table", False):
                continue
            grid = _table_grid(shape.table)
            if not grid:
                continue
            if pattern is None:
                candidates.append((grid, 0))
                continue
            for index, row in enumerate(grid):
                if pattern.search("|".join(row)):
                    candidates.append((grid, index))
                    break
    if not candidates:
        detail = f" with a row matching {header_match!r}" if header_match else ""
        raise SourceDataError(f"no visible slide table{detail}: {path.name}")
    if all_matching:
        chosen = candidates
    else:
        if table_index >= len(candidates):
            raise SourceDataError(f"{len(candidates)} candidate table(s), no table_index {table_index}: {path.name}")
        chosen = [candidates[table_index]]
    rows: list[dict[str, str]] = []
    for grid, header_row in chosen:
        keys: list[str] = []
        for position, cell in enumerate(grid[header_row]):
            key = cell or f"column_{position + 1}"
            keys.append(key if key not in keys else f"{key} ({position + 1})")
        if columns:
            keys = list(columns) + keys[len(columns):]
        for row in grid[header_row + 1:]:
            if not any(cell.strip() for cell in row):
                continue
            rows.append({keys[i]: cell for i, cell in enumerate(row)})
    return rows


class ReadTable(SourceCommand[PptxReadTableInput, PptxReadTableOutput]):
    """Read a slide table (located by its header row) as header-keyed rows."""

    name = "read_table"
    input_model = PptxReadTableInput
    output_model = PptxReadTableOutput

    def run(self, source_input: PptxReadTableInput, context: SourceContext) -> PptxReadTableOutput:
        return PptxReadTableOutput(read_pptx_table(
            context.resolve_path(source_input.path), source_input.header_match,
            source_input.table_index, source_input.all_matching, source_input.columns,
        ))

DESIGN_COMMANDS = (InspectShapes(), InspectFonts(), InspectBullets(), InspectFills(), ScanOmml(), ReadTable())
